import zipfile
from io import BytesIO

from docx import Document

from app.docgen import render
from app.images import normalize
from app.schemas import Submission


def make_sub(**overrides) -> Submission:
    data = {
        "category": "photo",
        "project_name": "工程一",
        "project_title": "河濱公園環境改善工程",
        "code": "JT-114-D-工審-026",
        "filename": "工審-026",
        "entries": [
            {"date": "2026-10-28", "item": "座椅鋼筋綁紮檢查", "location": "座椅", "desc": "主筋：#4-8支\n\n箍筋：#4@12.5cm"},
            {"date": None, "item": "A<B & C>D", "location": "", "desc": ""},
        ],
    }
    data.update(overrides)
    return Submission.model_validate(data)


def render_doc(sub, jpeg):
    photos = [normalize(jpeg((1600, 1200))), normalize(jpeg((600, 1600)))][: len(sub.entries)]
    raw = render(sub, photos)
    return raw, Document(BytesIO(raw))


def test_header_fields(jpeg):
    _, doc = render_doc(make_sub(), jpeg)
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "工程施工照片保存卡"
    assert "工程名稱：河濱公園環境改善工程" in texts
    assert "編號：JT-114-D-工審-026" in texts


def test_title_uses_category_name_when_no_doc_title(jpeg):
    _, doc = render_doc(make_sub(category="defect"), jpeg)
    assert doc.paragraphs[0].text == "缺失改善"


def test_one_row_per_entry_with_fields(jpeg):
    _, doc = render_doc(make_sub(), jpeg)
    table = doc.tables[0]
    assert len(table.rows) == 2  # 迴圈標記列已被拿掉
    first = table.rows[0].cells[0].text
    assert "日期：2026/10/28" in first
    assert "項目：座椅鋼筋綁紮檢查" in first
    assert "位置：座椅" in first
    # 說明的空行被去掉，其餘逐行保留
    assert "主筋：#4-8支\n箍筋：#4@12.5cm" in first
    second = table.rows[1].cells[0].text
    assert "日期：\n" in second or second.startswith("日期：")
    assert "項目：A<B & C>D" in second  # 特殊字元正確跳脫，不會弄壞 XML


def test_photos_embedded_and_fit_cell(jpeg):
    raw, doc = render_doc(make_sub(), jpeg)
    shapes = doc.inline_shapes
    assert len(shapes) == 2
    for shape in shapes:
        assert shape.width.pt <= 6554 / 20 - 16 + 0.5
        assert shape.height.pt <= 220 + 0.5
    with zipfile.ZipFile(BytesIO(raw)) as z:
        assert len([n for n in z.namelist() if n.startswith("word/media/")]) == 2


def test_no_template_tags_left(jpeg):
    raw, _ = render_doc(make_sub(), jpeg)
    with zipfile.ZipFile(BytesIO(raw)) as z:
        xml = z.read("word/document.xml").decode()
    for tag in ("{{", "}}", "{%", "%}"):
        assert tag not in xml


def test_jinja_in_user_input_is_not_executed(jpeg):
    sub = make_sub(project_title="{{ 7*7 }}", entries=[{"item": "{% for x in y %}", "desc": "{{ e.photo }}"}])
    _, doc = render_doc(sub, jpeg)
    assert "工程名稱：{{ 7*7 }}" in [p.text for p in doc.paragraphs]
    assert "{% for x in y %}" in doc.tables[0].rows[0].cells[0].text
