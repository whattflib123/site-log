"""產生 app/doc_templates/template.docx（docxtpl 範本）。

版面照舊版前端（index.html 的 buildDocument）：A4、標楷體 12pt、大標題 20pt 粗體底線，
兩欄表格每筆一列，左邊 日期/項目/位置/說明，右邊照片。

想手動改版面的話，直接用 Word 改 template.docx 也可以，Jinja 標籤（{{ }}、{%tr %}）要保留，
且每個標籤要在同一段同樣格式的文字裡（Word 有時會把一個標籤拆成好幾段，改完要實際產生一次確認）。

用法（在 form-app 目錄）：
  docker run --rm -v "$PWD":/srv -w /srv sitelog-form-app:local python tools/make_template.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

OUT = Path(__file__).resolve().parent.parent / "app" / "doc_templates" / "template.docx"

FONT = "標楷體"
PAGE_W, PAGE_H = 11906, 16838
MARGIN_L = MARGIN_R = 1134
CONTENT_W = PAGE_W - MARGIN_L - MARGIN_R  # 9638
COL0_W = round(CONTENT_W * 0.32)  # 文字欄
COL1_W = CONTENT_W - COL0_W  # 照片欄
ROW_H = 4680


def set_font(run_or_style, size=12, bold=None, underline=None):
    font = run_or_style.font
    font.name = FONT
    font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if underline is not None:
        font.underline = underline
    rpr = run_or_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def labelled(paragraph, label, value_tag):
    set_font(paragraph.add_run(label), bold=True)
    if value_tag:
        set_font(paragraph.add_run(value_tag), bold=False)


def cell_para(cell, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Twips(400)
    p.paragraph_format.space_after = Pt(0)
    return p


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def main():
    doc = Document()
    set_font(doc.styles["Normal"])

    section = doc.sections[0]
    section.page_width, section.page_height = Twips(PAGE_W), Twips(PAGE_H)
    section.top_margin, section.bottom_margin = Twips(340), Twips(663)
    section.left_margin, section.right_margin = Twips(MARGIN_L), Twips(MARGIN_R)

    title = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Twips(160)
    set_font(title.add_run("{{ doc_title }}"), size=20, bold=True, underline=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Twips(60)
    labelled(p, "工程名稱：", "{{ project_title }}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Twips(200)
    labelled(p, "編號：", "{{ code }}")

    # 第 0 列、第 2 列是 docxtpl 的迴圈標記，產生時會被拿掉
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    set_borders(table)
    # 欄寬要同時設在 tblGrid（columns）和每一格，不然 Word/LibreOffice 會平均分配
    table.columns[0].width, table.columns[1].width = Twips(COL0_W), Twips(COL1_W)
    for row in table.rows:
        row.cells[0].width, row.cells[1].width = Twips(COL0_W), Twips(COL1_W)
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table._tbl.tblPr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_W))
    tbl_w.set(qn("w:type"), "dxa")

    set_font(table.rows[0].cells[0].paragraphs[0].add_run("{%tr for e in entries %}"))
    set_font(table.rows[2].cells[0].paragraphs[0].add_run("{%tr endfor %}"))

    row = table.rows[1]
    row.height = Twips(ROW_H)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    text_cell, photo_cell = row.cells
    labelled(cell_para(text_cell, first=True), "日期：", "{{ e.date }}")
    labelled(cell_para(text_cell), "項目：", "{{ e.item }}")
    labelled(cell_para(text_cell), "位置：", "{{ e.location }}")
    labelled(cell_para(text_cell), "說明：", "")
    set_font(cell_para(text_cell).add_run("{{ e.desc }}"))

    photo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pp = photo_cell.paragraphs[0]
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run("{{ e.photo }}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
